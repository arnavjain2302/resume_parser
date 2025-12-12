from docx import Document
import json
import re
from google import genai
import time
import io
from markdown_pdf import MarkdownPdf, Section


def extract_docx_text(file):
    doc = Document(file)
    paras = [p.text.strip() for p in doc.paragraphs]
    final_paras = []
    for p in paras:
        if p:
            final_paras.append(p)
        elif not final_paras or final_paras[-1] != "":
            final_paras.append("")
    return "\n".join(final_paras).strip()
 
def basic_clean(text):
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(re.sub(r"[ \t]{2,}", " ", line) for line in text.splitlines())
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    return text

SCHEMA = {
  "name": "",
  "email": "",
  "phone": "",
  
  "education": [
    {
      "degree": "",
      "institution": "",
      "start_date": "",
      "end_date": "",
      "cgpa/percentage": ""
    }
  ],
  
  "experience": [
    {
      "job_title": "",
      "company": "",
      "location": "",
      "start_date": "",
      "end_date": "",
      "description": ""
    }
  ],
  
  "projects": [
    {
      "name": "",
      "supervisor": "",
      "location": " ",
      "description": ""
    }
  ],
  "skills": [],
  "publications": [
      {
          "name": "",
          "conference": "",
          "first author": "",
          "second author(s)": "",
      }
  ]

}

OPTIONAL_ATTRIBUTES = {
    "patents": {
        "patents": [
            {
                "title": "",
                "year": "",
                "description": ""
            }
        ]
    },
    "awards": {
        "awards": [
            {
                "name": "",
                "year": "",
                "issuer": "",
                "description": ""
            }
        ]
    },
    "scholarships": {
        "scholarships": [
            {
                "name": "",
                "year": "",
                "issuer": "",
                "amount": "",
                "description": ""
            }
        ]
    }
}

def merge_schemas(base_schema: dict, extras_keys: list) -> dict:
    
    import copy
    final = copy.deepcopy(base_schema)
    for k in extras_keys:
        tmpl = OPTIONAL_ATTRIBUTES.get(k)
        if not tmpl:
            continue
        for key, val in tmpl.items():
            if key not in final:
                final[key] = val
    return final


def build_prompt(text:str , schema : dict):
    schema_str = json.dumps(schema, indent=2, ensure_ascii=False)
    return f"""
You are a resume parser.

Convert the resume text below into a single JSON object that EXACTLY matches this schema:

{schema_str}

Rules:
- Output ONLY valid JSON. No extra text, no backticks.
- Use "" or [] when information is missing.
- For education/experience/patents/awards/scholarships, return lists of objects exactly in this structure.
- Infer values only from the resume text.

----- RESUME TEXT -----
{text}
----- END -----
""".strip()

def extract_json(s):
    start = s.find('{')
    if start == -1:
        raise ValueError("No JSON found in model output.")
    depth = 0
    for i in range(start, len(s)):
        if s[i] == '{':
            depth += 1
        elif s[i] == '}':
            depth -= 1
            if depth == 0:
                return s[start:i+1]

def parse_resume(file, api_key, schema_1):
    client = genai.Client(api_key=api_key)
    
    raw = extract_docx_text(file)
    cleaned = basic_clean(raw)

    prompt = build_prompt(cleaned , schema_1)

    time.sleep(1)

    response = client.models.generate_content(
        model="gemini-2.5-flash-lite",
        contents=prompt
    )

    raw_output = response.text
    json_str = extract_json(raw_output)
    return json.loads(json_str)

def llm_strict_format(api_key, parsed_json):
    client = genai.Client(api_key=api_key)
    import json
    
    prompt = f"""
    SYSTEM: You are a formatting assistant. You MUST NOT invent new facts or add information not present in the JSON.
    You MAY rephrase, clean, and polish the existing text to improve readability and recruiter-readiness.
    Output only Markdown. Do not output JSON or explanation.

    USER: Convert the following JSON into a polished Markdown resume. Follow these exact rules:
    1) Use headings (# Name, ## Experience, ## Projects, ## Skills, etc.).
    2) For each experience/project, produce a bold job title/company line, a short italic date/location line, and then 1-3 concise bullet points that **paraphrase only the description fields provided**. Do NOT add new responsibilities, awards, numbers, or claims not present in the JSON.
    3) Keep all original factual values (company, dates, locations, titles) unchanged.
    4) When paraphrasing, keep the meaning and facts exact — you may change phrasing, grammar, and sentence order only.
    5) If a description is long, compress into 1-3 concise bullets using only the existing info.
    6) If a field is empty, skip it.
    7) Output only Markdown.
    Now convert this JSON (do NOT invent anything):
    {json.dumps(parsed_json, indent=2, ensure_ascii=False)}
    """

    time.sleep(1)

    resp = client.models.generate_content(
        model="gemini-2.5-flash-lite",
        contents=prompt
    )
    text = resp.text if hasattr(resp, "text") else str(resp)
    return text.strip()


def text_to_docx(text: str):
    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def text_to_pdf(text: str):
    pdf = MarkdownPdf()
    pdf.add_section(Section(text))
    
    buffer = io.BytesIO()
    pdf.save_bytes(buffer)
    buffer.seek(0)
    return buffer

